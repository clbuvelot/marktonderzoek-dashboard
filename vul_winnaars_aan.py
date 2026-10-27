"""
vul_winnaars_aan.py
-------------------
Vult winnaarsnamen aan in de bestaande dataset, uit de TenderNed bulk Excel.

Waarom een apart script: de TNS API geeft geen winnaar terug. Niet in de
lijst-respons en niet op het detail-endpoint — daar staat alleen isGegund.
De naam van de gegunde onderneming staat uitsluitend in de bulk Excel die
TenderNed periodiek publiceert.

Verschil met koppel_winnaar.py: dat script bouwt geclassificeerd_met_winnaar
opnieuw op vanuit geclassificeerd.json en gooit daarmee alles weg wat update.py
sindsdien heeft toegevoegd. Dit script werkt de bestaande dataset bij en laat
records die het niet kent met rust.

Gebruik:
    python vul_winnaars_aan.py
    python vul_winnaars_aan.py --excel pad/naar/dataset.xlsx   # zelf aanleveren
    python vul_winnaars_aan.py --dry-run

Leest/schrijft:
    data/geclassificeerd_met_winnaar.json   (bijgewerkt op z'n plek)
    data/concurrentiekaart.csv              (opnieuw opgebouwd)
"""

import argparse
import json
import os
import re
import tempfile
import urllib.request

import openpyxl

from koppel_winnaar import bouw_concurrentiekaart, sla_op_csv

DATASET = "data/geclassificeerd_met_winnaar.json"
OUTPUT_DIR = "data"
DATASETS_PAGINA = "https://www.tenderned.nl/cms/aanbesteden-in-cijfers/datasets-aanbestedingen"
SHEET = "OpenData sheet"

KOLOMMEN = {
    "id": "ID publicatie",
    "kenmerk": "Tenderned kenmerk",
    "winnaar": "Naam gegunde onderneming",
    "datum": "Datum gunning",
    "waarde": "Definitieve waarde - bedrag",
}


# ── Excel vinden en ophalen ───────────────────────────────────────────────────

def zoek_nieuwste_excel_url() -> str:
    """
    Pakt van de TenderNed-datasetpagina het bestand met de laatste einddatum.
    De bestandsnaam bevat die datum, bijvoorbeeld:
    Dataset_Tenderned-compleet-2021-01-01-2026-10-01.xlsx
    """
    print("[1/5] Nieuwste bulk Excel zoeken op tenderned.nl...")
    with urllib.request.urlopen(DATASETS_PAGINA, timeout=60) as resp:
        html = resp.read().decode("utf-8", errors="replace")

    kandidaten = []
    for url in set(re.findall(r'href="([^"]+\.xlsx)"', html, re.I)):
        datums = re.findall(r"(\d{4}-\d{2}-\d{2})", url)
        if datums:
            kandidaten.append((max(datums), url))

    if not kandidaten:
        raise RuntimeError("Geen .xlsx-bestanden gevonden op de datasetpagina")

    einddatum, url = max(kandidaten)
    if not url.startswith("http"):
        url = "https://www.tenderned.nl" + url
    print(f"  gevonden, loopt tot {einddatum}")
    return url


def download_excel(url: str) -> str:
    doel = os.path.join(tempfile.gettempdir(), "tenderned_bulk.xlsx")
    if os.path.exists(doel):
        print(f"[2/5] Excel al gedownload: {doel}")
        return doel
    print(f"[2/5] Downloaden (ongeveer 90 MB, dit duurt even)...")
    urllib.request.urlretrieve(url, doel)
    print(f"  {round(os.path.getsize(doel) / 1e6, 1)} MB opgeslagen")
    return doel


# ── Winnaars uitlezen ─────────────────────────────────────────────────────────

def bouw_lookup(pad: str) -> tuple[dict, dict]:
    """
    Leest de Excel streamend uit en bouwt twee opzoektabellen: op kenmerk en op
    publicatie-id. Kenmerk is de betere sleutel — die overleeft het feit dat een
    uitvraag en de bijbehorende gunning verschillende publicatie-id's hebben.
    """
    print("[3/5] Winnaars uitlezen uit de Excel...")
    wb = openpyxl.load_workbook(pad, read_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"Sheet {SHEET!r} ontbreekt; gevonden: {wb.sheetnames}")

    rijen = wb[SHEET].iter_rows(values_only=True)
    kop = list(next(rijen))
    ontbreekt = [k for k in KOLOMMEN.values() if k not in kop]
    if ontbreekt:
        raise RuntimeError(f"Kolommen ontbreken in de Excel: {ontbreekt}")
    ix = {naam: kop.index(kolom) for naam, kolom in KOLOMMEN.items()}

    op_kenmerk, op_id = {}, {}
    gelezen = 0
    for rij in rijen:
        gelezen += 1
        winnaar = rij[ix["winnaar"]]
        if not winnaar or not str(winnaar).strip():
            continue
        waarde = rij[ix["waarde"]]
        try:
            waarde = float(waarde) if waarde is not None else None
        except (TypeError, ValueError):
            waarde = None
        info = {
            "winnaar": str(winnaar).strip(),
            "waarde_definitief": waarde,
            "datum_gunning": str(rij[ix["datum"]] or "").strip(),
        }
        kenmerk = str(rij[ix["kenmerk"]] or "").strip()
        pub_id = str(rij[ix["id"]] or "").strip()

        # Percelen: meerdere winnaars onder hetzelfde kenmerk.
        for tabel, sleutel in ((op_kenmerk, kenmerk), (op_id, pub_id)):
            if not sleutel:
                continue
            if sleutel not in tabel:
                tabel[sleutel] = {**info, "extra_winnaars": []}
            elif info["winnaar"] != tabel[sleutel]["winnaar"]:
                if info["winnaar"] not in tabel[sleutel]["extra_winnaars"]:
                    tabel[sleutel]["extra_winnaars"].append(info["winnaar"])

    wb.close()
    print(f"  {gelezen} rijen gelezen, {len(op_kenmerk)} kenmerken met winnaar")
    return op_kenmerk, op_id


# ── Koppelen ──────────────────────────────────────────────────────────────────

def vul_aan(records: list[dict], op_kenmerk: dict, op_id: dict) -> dict:
    """Vult winnaarsvelden aan. Bestaande winnaars blijven staan."""
    tellers = {"nieuw": 0, "al_bekend": 0, "geen_match": 0}

    for r in records:
        if (r.get("winnaar") or "").strip():
            tellers["al_bekend"] += 1
            continue

        kenmerk = str(r.get("tenderned_kenmerk") or "").strip()
        info = op_kenmerk.get(kenmerk)
        if not info:
            for sleutel in (r.get("gunning_publicatie_id"), r.get("id")):
                info = op_id.get(str(sleutel or "").strip())
                if info:
                    break
        if not info:
            tellers["geen_match"] += 1
            continue

        r["winnaar"] = info["winnaar"]
        r["waarde_definitief"] = info["waarde_definitief"]
        # Het dashboard leest gunningswaarde, koppel_winnaar.py waarde_definitief.
        r["gunningswaarde"] = info["waarde_definitief"]
        r["datum_gunning"] = info["datum_gunning"]
        if info["extra_winnaars"]:
            r["extra_winnaars"] = info["extra_winnaars"]
        if r.get("publicatie_type") != "gegund":
            r["publicatie_type"] = "gegund"
        tellers["nieuw"] += 1

    return tellers


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", help="Pad naar een al gedownloade bulk Excel")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(DATASET):
        print(f"[FOUT] {DATASET} niet gevonden.")
        return

    pad = args.excel or download_excel(zoek_nieuwste_excel_url())
    op_kenmerk, op_id = bouw_lookup(pad)

    print("[4/5] Dataset bijwerken...")
    with open(DATASET, encoding="utf-8") as f:
        records = json.load(f)
    print(f"  {len(records)} records geladen")

    tellers = vul_aan(records, op_kenmerk, op_id)
    print(f"  {tellers['nieuw']} nieuwe winnaars, "
          f"{tellers['al_bekend']} al bekend, "
          f"{tellers['geen_match']} zonder match")

    if args.dry_run:
        print("\n[DRY RUN] Niets opgeslagen.")
        return

    if not tellers["nieuw"]:
        print("\nGeen nieuwe winnaars. Niets gewijzigd.")
        return

    print("[5/5] Opslaan...")
    with open(DATASET, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    print(f"  → {DATASET}")

    kaart = bouw_concurrentiekaart(records)
    sla_op_csv(
        kaart,
        f"{OUTPUT_DIR}/concurrentiekaart.csv",
        ["bureau", "aantal_gewonnen", "waarvan_kwantitatief", "kwantitatief_pct",
         "totaalwaarde_euro", "sterkste_thema", "sterkste_methodiek"],
    )

    met_winnaar = sum(1 for r in records if (r.get("winnaar") or "").strip())
    print(f"\n=== Klaar ===")
    print(f"Records met winnaar : {met_winnaar} van {len(records)}")
    print(f"Bureaus in kaart    : {len(kaart)}")


if __name__ == "__main__":
    main()
